#!/usr/bin/env python3
"""Extract JeS headline totals and staff lines from Swansea Finance budget xlsx."""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.stderr.write("Install openpyxl: pip install openpyxl\n")
    sys.exit(1)


def cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def money_from_row(cells: list[str]) -> list[float]:
    out: list[float] = []
    for c in cells:
        c = c.replace("£", "").replace(",", "").strip()
        if re.fullmatch(r"\d+(\.\d+)?", c):
            out.append(float(c))
    return out


def scan_workbook(path: Path) -> None:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    print(f"# Source: {path.name}\n")
    project_total: tuple[float, float] | None = None
    staff_rows: list[tuple[str, str]] = []

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            cells = [cell_str(v) for v in row]
            line = " ".join(c for c in cells if c)
            lower = line.lower()
            if "project total" in lower or "university total" in lower:
                nums = money_from_row(cells)
                if len(nums) >= 2:
                    project_total = (nums[-2], nums[-1])
            if any(
                k in lower
                for k in (
                    "research associate",
                    "project lead",
                    "principal investigator",
                    "co-investigator",
                    "project co-lead",
                    "villamizar",
                    "brown",
                )
            ):
                nums = money_from_row(cells)
                if nums and any(n > 1000 for n in nums):
                    cost = max(n for n in nums if n > 1000)
                    staff_rows.append((sheet, f"{line} | £{cost:,.0f}"))

    if project_total:
        fec, epsrc = project_total
        print(f"**Total fEC:** £{fec:,.0f}")
        print(f"**EPSRC 80%:** £{epsrc:,.0f}")
        print(f"**Headroom to £100k fEC:** £{max(0, 100_000 - fec):,.0f}\n")
    else:
        print("Could not find Project/University total row; check sheet layout.\n")

    if staff_rows:
        print("## Candidate staff lines (verify against Finance sheet)\n")
        for sheet, text in staff_rows:
            print(f"- [{sheet}] {text}")
    else:
        print("No staff cost rows detected automatically; copy from Finance sheet manually.")

    wb.close()


def main() -> None:
    if len(sys.argv) < 2:
        default = Path(__file__).parent / "Budget 02 10 26 IPA 799 v2.xlsx"
        path = default if default.is_file() else None
        if path is None:
            sys.stderr.write(
                "Usage: extract_jes_totals.py <Budget 02 10 26 IPA 799 v2.xlsx>\n"
            )
            sys.exit(1)
    else:
        path = Path(sys.argv[1])
    if not path.is_file():
        sys.stderr.write(f"File not found: {path}\n")
        sys.exit(1)
    scan_workbook(path)


if __name__ == "__main__":
    main()
